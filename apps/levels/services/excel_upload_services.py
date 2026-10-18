from __future__ import annotations

import re
import unicodedata
from zipfile import BadZipFile
from typing import Any

from django import forms
from openai import OpenAIError
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from ..ai.factory import get_ai_client
from ..forms import ClassPlanningCreateForm, PerformanceLevelTemplateCreateForm
from ..models import PerformanceLevelTemplate
from ..models.performance_level_template_model import (
    AREAS_OPCIONES,
    GRADOS_OPCIONES,
    PERIODO_ACADEMICO_OPCIONES,
)


class ExcelUploadError(Exception):
    """
    Error de negocio al procesar el archivo Excel.
    """


# Campos de ClassPlanningCreateForm que la IA debe identificar en el Excel
CLASS_PLANNING_EXCEL_FIELDS = (
    "topic",
    "subtopic",
    "class_objective",
    "duration_minutes",
    "prompt",
    "methodology",
    "resources",
)


# Límite de texto del Excel que se envía a la IA (evita prompts demasiado grandes)
EXCEL_MAX_TEXT_CHARS = 30000


# Nombres alternativos del área que se aceptan en el Excel (normalizados: minúsculas y sin tildes)
AREA_ALIASES = {
    "lenguaje": "Español / Castellano",
    "lengua castellana": "Español / Castellano",
    "etica": "Ética y Valores",
    "e. fisica": "Educación Física",
}


class ExcelUploadService:

    @staticmethod
    def create_performance_level_from_excel_service(*, excel_file, user) -> PerformanceLevelTemplate:
        """
        Lee el archivo Excel con IA y crea directamente la plantilla para el usuario
        (usado por el administrador al cargar un Excel a nombre de un docente).
        """
        data = ExcelUploadService.extract_performance_level_data_from_excel_service(excel_file=excel_file)

        # Se reutiliza el formulario de creación para validar los datos
        form = PerformanceLevelTemplateCreateForm(data=data)
        if not form.is_valid():
            raise ExcelUploadError(ExcelUploadService._format_errors(form))

        performance_level = form.save(commit=False)
        performance_level.user = user
        performance_level.save()

        return performance_level

    @staticmethod
    def extract_performance_level_data_from_excel_service(*, excel_file) -> dict[str, Any]:
        """
        Lee el archivo Excel y la IA identifica el valor de cada campo de
        PerformanceLevelTemplateCreateForm para precargar el formulario.
        No crea ningún registro.
        """
        form = PerformanceLevelTemplateCreateForm()
        data = ExcelUploadService._extract_data_with_ai(excel_file, form, form.fields.keys())

        if not any(data.values()):
            raise ExcelUploadError("No se encontró información de la plantilla en el archivo.")

        return ExcelUploadService._normalize_choices(data)

    @staticmethod
    def extract_class_planning_data_from_excel_service(*, excel_file) -> dict[str, Any]:
        """
        Lee el archivo Excel y la IA identifica los campos definidos en
        CLASS_PLANNING_EXCEL_FIELDS para precargar el formulario de
        planeación de clase. No crea ningún registro.
        """
        data = ExcelUploadService._extract_data_with_ai(
            excel_file, ClassPlanningCreateForm(), CLASS_PLANNING_EXCEL_FIELDS
        )

        if not any(data.values()):
            raise ExcelUploadError("No se encontró información de la planeación de clase en el archivo.")

        # La duración puede venir como "60", "60.0" o "60 min"
        match = re.match(r"\d+", data["duration_minutes"])
        data["duration_minutes"] = match.group() if match else ""

        return data

    @staticmethod
    def _extract_data_with_ai(excel_file, form, field_names) -> dict[str, str]:
        """
        Envía el contenido del Excel a la IA junto con la descripción de los
        campos del formulario y devuelve {campo: valor}.
        """
        excel_text = ExcelUploadService._read_workbook_as_text(excel_file)
        if not excel_text:
            raise ExcelUploadError("El archivo Excel está vacío.")

        fields_schema = ExcelUploadService._build_fields_schema(form, field_names)

        try:
            return get_ai_client().extract_form_fields_from_excel(excel_text, fields_schema)
        except (OpenAIError, ValueError) as exc:
            raise ExcelUploadError(
                "No se pudo analizar el archivo con IA. Intente de nuevo en unos minutos."
            ) from exc

    @staticmethod
    def _read_workbook_as_text(excel_file) -> str:
        """
        Convierte todas las hojas del Excel en texto, una línea por fila con
        las celdas que tienen contenido (por ejemplo "Fila 2: A=Matemáticas | B=5").
        """
        try:
            workbook = load_workbook(excel_file, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as exc:
            raise ExcelUploadError("No se pudo leer el archivo. Verifique que sea un Excel válido.") from exc

        try:
            lines = []
            for sheet in workbook.worksheets:
                sheet_lines = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    cells = [
                        f"{get_column_letter(column)}={str(value).strip()}"
                        for column, value in enumerate(row, start=1)
                        if value is not None and str(value).strip()
                    ]
                    if cells:
                        sheet_lines.append(f"Fila {row_number}: {' | '.join(cells)}")

                if sheet_lines:
                    lines.append(f"### Hoja: {sheet.title}")
                    lines.extend(sheet_lines)
        finally:
            workbook.close()

        return "\n".join(lines)[:EXCEL_MAX_TEXT_CHARS]

    @staticmethod
    def _build_fields_schema(form, field_names) -> list[dict[str, Any]]:
        """
        Describe los campos del formulario (etiqueta, ayuda, tipo y opciones)
        para que la IA sepa qué buscar en el Excel.
        """
        schema = []
        for name in field_names:
            field = form.fields[name]
            field_schema = {
                "name": name,
                "label": str(field.label or name),
                "help_text": str(field.help_text or ""),
                "type": "number" if isinstance(field, (forms.IntegerField, forms.DecimalField)) else "text",
            }
            if isinstance(field, forms.ChoiceField) and not isinstance(field, forms.ModelChoiceField):
                field_schema["choices"] = [(str(key), str(label)) for key, label in field.choices if key != ""]
            schema.append(field_schema)
        return schema

    @staticmethod
    def _extract_data(excel_file, cell_mapping) -> dict[str, Any]:
        """
        Lee las celdas fijas definidas en cell_mapping (usado por Prompt Lab).
        """
        try:
            workbook = load_workbook(excel_file, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as exc:
            raise ExcelUploadError("No se pudo leer el archivo. Verifique que sea un Excel válido.") from exc

        try:
            sheet = workbook.active
            data = {}
            for cell, field in cell_mapping.items():
                value = sheet[cell].value
                data[field] = "" if value is None else str(value).strip()
            return data
        finally:
            workbook.close()

    @staticmethod
    def _normalize_choices(data: dict[str, Any]) -> dict[str, Any]:
        """
        Ajusta los valores del Excel a las opciones del modelo
        (ignora mayúsculas y tildes; el grado acepta número o nombre).
        """
        area = AREA_ALIASES.get(ExcelUploadService._normalize_text(data["area"]), data["area"])
        data["area"] = ExcelUploadService._match_choice(area, AREAS_OPCIONES)
        data["academic_period"] = ExcelUploadService._match_choice(
            data["academic_period"], PERIODO_ACADEMICO_OPCIONES
        )

        grade = data["grade"]
        if grade.endswith(".0"):
            grade = grade[:-2]
        data["grade"] = ExcelUploadService._match_choice(grade, GRADOS_OPCIONES)

        return data

    @staticmethod
    def _match_choice(value: str, choices: list[tuple[str, str]]) -> str:
        normalized_value = ExcelUploadService._normalize_text(value)
        for key, label in choices:
            # "Español / Castellano" también acepta "Español" o "Castellano"
            options = {key, label, *key.split("/"), *label.split("/")}
            if normalized_value in {ExcelUploadService._normalize_text(option) for option in options}:
                return key
        return value

    @staticmethod
    def _normalize_text(value: str) -> str:
        value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
        return " ".join(value.lower().split())

    @staticmethod
    def _format_errors(form) -> str:
        errors = []
        for field, field_errors in form.errors.items():
            label = form.fields[field].label if field in form.fields else field
            errors.append(f"{label}: {' '.join(field_errors)}")
        return " | ".join(errors)
